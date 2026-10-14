"""Source opening: authorised media, PDF page images and slide views for citations.

<video>/<img>/<iframe> elements can't send an Authorization header, so the client first asks for a
short-lived, document-scoped media token and uses it as a query parameter.
"""
import io
import re
import uuid
from datetime import datetime, timedelta, timezone

import jwt
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse, Response
from sqlalchemy.orm import Session

from app.api.deps import current_user, owned_document
from app.core.config import get_settings
from app.database.session import get_db
from app.models.course import Document
from app.models.user import User

router = APIRouter(prefix="/api", tags=["media"])
MEDIA_TOKEN_MINUTES = 15


def _mint(user_id: uuid.UUID, doc_id: uuid.UUID) -> str:
    s = get_settings()
    exp = datetime.now(timezone.utc) + timedelta(minutes=MEDIA_TOKEN_MINUTES)
    return jwt.encode({"sub": str(user_id), "doc": str(doc_id), "typ": "media", "exp": exp}, s.jwt_secret, algorithm=s.jwt_algorithm)


def _doc_from_media_token(doc_id: uuid.UUID, token: str, db: Session) -> Document:
    s = get_settings()
    try:
        data = jwt.decode(token, s.jwt_secret, algorithms=[s.jwt_algorithm])
    except jwt.PyJWTError:
        raise HTTPException(401, "This link has expired. Reopen the source.")
    if data.get("typ") != "media" or data.get("doc") != str(doc_id):
        raise HTTPException(401, "Invalid media link.")
    user = db.get(User, uuid.UUID(data["sub"]))
    doc = db.get(Document, doc_id)
    if user is None or doc is None or doc.course.owner_id != user.id:  # ownership re-checked at use time
        raise HTTPException(404, "Document not found.")
    return doc


@router.get("/documents/{document_id}/media-token")
def media_token(document_id: uuid.UUID, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = owned_document(document_id, user, db)
    return {"token": _mint(user.id, doc.id), "expires_in": MEDIA_TOKEN_MINUTES * 60, "source_type": doc.source_type,
            "page_count": doc.page_count, "slide_count": doc.slide_count}


@router.get("/media/documents/{document_id}")
def media_file(document_id: uuid.UUID, t: str = Query(...), db: Session = Depends(get_db)):
    doc = _doc_from_media_token(document_id, t, db)
    path = get_settings().upload_root / doc.stored_path
    if not path.is_file():
        raise HTTPException(404, "The stored file is missing.")
    return FileResponse(path, media_type=doc.mime_type, filename=doc.original_filename, content_disposition_type="inline")  # supports Range


@router.get("/media/documents/{document_id}/pages/{page}.png")
def pdf_page_image(document_id: uuid.UUID, page: int, t: str = Query(...), db: Session = Depends(get_db)):
    doc = _doc_from_media_token(document_id, t, db)
    if doc.source_type != "pdf" or not doc.page_count or not (1 <= page <= doc.page_count):
        raise HTTPException(404, "Page not found.")
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(get_settings().upload_root / doc.stored_path))
    try:
        buf = io.BytesIO()
        pdf[page - 1].render(scale=1.6).to_pil().save(buf, "PNG")
    finally:
        pdf.close()
    return Response(buf.getvalue(), media_type="image/png", headers={"Cache-Control": "private, max-age=600"})


@router.get("/media/documents/{document_id}/figures/{figure_ref}")
def figure_image(document_id: uuid.UUID, figure_ref: str, t: str = Query(...), db: Session = Depends(get_db)):
    """Return an authorized thumbnail for a figure reference stored on an indexed chunk."""
    doc = _doc_from_media_token(document_id, t, db)
    path = get_settings().upload_root / doc.stored_path
    if not path.is_file():
        raise HTTPException(404, "The stored file is missing.")

    from PIL import Image, ImageOps

    image = None
    if doc.source_type == "image" and figure_ref == doc.original_filename:
        return FileResponse(path, media_type=doc.mime_type, headers={"Cache-Control": "private, max-age=600"})
    if doc.source_type == "pdf":
        match = re.fullmatch(r"page(\d+)-img(\d+)", figure_ref)
        if not match:
            raise HTTPException(404, "Figure not found.")
        page_no, image_no = map(int, match.groups())
        import pypdfium2 as pdfium

        from app.services.ingestion.extract import _image_object_to_pil, pdf_figure_objects

        pdf = pdfium.PdfDocument(str(path))
        try:
            if not (1 <= page_no <= len(pdf)):
                raise HTTPException(404, "Figure not found.")
            image_obj = next((obj for index, obj in pdf_figure_objects(pdf[page_no - 1]) if index == image_no), None)
            if image_obj is None:
                raise HTTPException(404, "Figure not found.")
            image = _image_object_to_pil(image_obj)
        finally:
            pdf.close()
    elif doc.source_type == "pptx":
        match = re.fullmatch(r"slide(\d+)-img(\d+)", figure_ref)
        if not match:
            raise HTTPException(404, "Figure not found.")
        slide_no, image_no = map(int, match.groups())
        from pptx import Presentation

        from app.services.ingestion.extract import _iter_pictures, pptx_picture_is_figure

        prs = Presentation(str(path))
        if not (1 <= slide_no <= len(prs.slides)):
            raise HTTPException(404, "Figure not found.")
        slide = prs.slides[slide_no - 1]
        picture = next((pic for index, pic in enumerate(_iter_pictures(slide.shapes), start=1)
                        if index == image_no and pptx_picture_is_figure(pic, prs.slide_width, prs.slide_height)), None)
        if picture is None:
            raise HTTPException(404, "Figure not found.")
        try:
            image = Image.open(io.BytesIO(picture.image.blob)).copy()
        except Exception as exc:
            raise HTTPException(404, "Figure could not be decoded.") from exc
    else:
        raise HTTPException(404, "Figure not found.")

    image = ImageOps.contain(image.convert("RGB"), (1200, 900))
    buf = io.BytesIO()
    image.save(buf, "PNG", optimize=True)
    return Response(buf.getvalue(), media_type="image/png", headers={"Cache-Control": "private, max-age=600"})


@router.get("/documents/{document_id}/slides/{slide}")
def slide_view(document_id: uuid.UUID, slide: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """The slide's own extracted content (title, body, notes). Not a pixel render of the slide."""
    doc = owned_document(document_id, user, db)
    if doc.source_type != "pptx" or not doc.slide_count or not (1 <= slide <= doc.slide_count):
        raise HTTPException(404, "Slide not found.")
    from pptx import Presentation

    from app.services.ingestion.extract import _shape_text

    s = list(Presentation(str(get_settings().upload_root / doc.stored_path)).slides)[slide - 1]
    title = s.shapes.title.text_frame.text.strip() if s.shapes.title is not None and s.shapes.title.has_text_frame else None
    body: list[str] = []
    for shape in s.shapes:
        for part in _shape_text(shape):
            if part != title:
                body.append(part)
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    return {"slide": slide, "of": doc.slide_count, "title": title, "text": body, "notes": notes or None}
