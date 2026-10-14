"""Builders for small, real sample files used by tests (generated, not stored)."""
import io
import shutil
import subprocess
from pathlib import Path


def make_text_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF whose pages each carry the given text in a real text layer."""
    objs: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        stream = f"BT /F1 14 Tf 72 700 Td ({text}) Tj ET".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_pdf_with_figure(text: str = "A diagram compares the efficiency of two heat engines.") -> bytes:
    """Build a text PDF with one substantial embedded JPEG at a known page location."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (480, 320), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 450, 290), outline="navy", width=6)
    draw.line((80, 240, 400, 80), fill="red", width=8)
    draw.text((80, 250), "Heat engine diagram", fill="black")
    image_buf = io.BytesIO()
    image.save(image_buf, "JPEG")
    jpg = image_buf.getvalue()

    content = (f"BT /F1 14 Tf 72 700 Td ({text}) Tj ET\n"
               "q 180 0 0 120 120 480 cm /Im1 Do Q\n").encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> /XObject << /Im1 6 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /XObject /Subtype /Image /Width 480 /Height 320 /ColorSpace /DeviceRGB "
        b"/BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n" % len(jpg) + jpg + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for object_id, obj in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{object_id} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def text_image(text: str | None, size=(1000, 500)):
    """White image with `text` drawn large and dark so Tesseract can read it. None => blank."""
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", size, "white")
    if text:
        font = ImageFont.load_default(size=44)
        d = ImageDraw.Draw(img)
        y = 40
        for line in text.split("\n"):
            d.text((40, y), line, fill="black", font=font)
            y += 80
    return img


def make_scanned_pdf(text: str | None = "Entropy always increases") -> bytes:
    """A PDF containing only a raster image (no text layer), like a scan."""
    buf = io.BytesIO()
    text_image(text).save(buf, "PDF", resolution=100)
    return buf.getvalue()


def make_pptx(slide_texts: list[str], notes: dict[int, str] | None = None, table_on: int | None = None,
              pictures: dict[int, list[tuple[bytes, float]]] | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, t in enumerate(slide_texts, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = t
        slide.placeholders[1].text = f"Body for {t}"
        if notes and i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
        if table_on == i:
            tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
            for r, row in enumerate([["Engine", "Efficiency"], ["Carnot", "Maximum"]]):
                for c, v in enumerate(row):
                    tbl.cell(r, c).text = v
        for data, inches in (pictures or {}).get(i, []):
            slide.shapes.add_picture(io.BytesIO(data), Inches(1), Inches(2), width=Inches(inches), height=Inches(inches * 0.7))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_png(text: str | None = None) -> bytes:
    """Small solid PNG, or a text image when `text` is given."""
    from PIL import Image

    buf = io.BytesIO()
    (text_image(text) if text else Image.new("RGB", (64, 48), "navy")).save(buf, "PNG")
    return buf.getvalue()


def make_mp4(tmp_dir: Path, seconds: int = 2, audio: bool = True) -> bytes | None:
    """Generate a tiny real video with FFmpeg. Returns None if FFmpeg is unavailable."""
    if shutil.which("ffmpeg") is None:
        return None
    out = tmp_dir / ("with_audio.mp4" if audio else "silent.mp4")
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"color=c=blue:s=160x120:d={seconds}"]
    if audio:
        cmd += ["-f", "lavfi", "-i", f"sine=frequency=440:duration={seconds}", "-shortest"]
    cmd += ["-pix_fmt", "yuv420p", str(out)]
    subprocess.run(cmd, check=True)
    return out.read_bytes()


VTT = b"WEBVTT\n\n00:00:01.000 --> 00:00:04.000\nHeat flows from hot to cold.\n\n00:00:04.000 --> 00:00:08.500\nEntropy never decreases in an isolated system.\n"
